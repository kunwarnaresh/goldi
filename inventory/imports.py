"""Excel import: Upload -> Validate -> Preview -> Error report -> (correct, re-upload) -> Import -> Success report.

Uploaded rows are staged in ImportBatch.rows and never written to master/stock tables until a batch
validates cleanly; the import itself re-validates and runs in one transaction through the posting engine.
"""
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db import transaction
from openpyxl import Workbook, load_workbook

from .engine import InventoryError
from .models import Bin, ImportBatch, JewelleryUnit, Location, SKU
from .services import create_transfer_order, find_unit, post_receipt, register_unit
from .tenancy import allowed_locations

TEMPLATES = {
    'opening_inventory': ['sku_code', 'location_code', 'bin_code', 'quantity', 'unit_cost', 'barcode', 'serial_no', 'huid',
                          'gross_weight', 'stone_weight', 'other_weight', 'purity', 'certificate_no'],
    'locations': ['code', 'name', 'location_type', 'city', 'state', 'pin', 'gstin', 'is_warehouse', 'allow_pos'],
    'transfer_orders': ['transfer_ref', 'from_location', 'to_location', 'sku_code', 'barcode', 'quantity', 'from_bin', 'to_bin'],
}
REQUIRED = {
    'opening_inventory': ['sku_code', 'location_code'],
    'locations': ['code', 'name', 'location_type'],
    'transfer_orders': ['transfer_ref', 'from_location', 'to_location'],
}


def template_workbook(import_type):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = import_type
    sheet.append(TEMPLATES[import_type])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def error_workbook(batch):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['row', 'column', 'message'])
    for error in batch.errors:
        sheet.append([error['row'], error.get('column', ''), error['message']])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def upload(actor, import_type, file):
    if import_type not in TEMPLATES:
        raise InventoryError('Unknown import type.')
    try:
        sheet = load_workbook(file, read_only=True, data_only=True).active
    except Exception:
        raise InventoryError('The file is not a readable .xlsx workbook.')
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise InventoryError('The workbook is empty.')
    header = [str(h or '').strip().lower() for h in rows[0]]
    missing = [c for c in REQUIRED[import_type] if c not in header]
    if missing:
        raise InventoryError(f'Missing column(s): {", ".join(missing)}. Download the template.')
    data = []
    for values in rows[1:]:
        if not any(v not in (None, '') for v in values):
            continue
        data.append({header[i]: ('' if v is None else str(v).strip()) for i, v in enumerate(values) if i < len(header) and header[i]})
    batch = ImportBatch.objects.create(tenant=actor.tenant, import_type=import_type, file_name=getattr(file, 'name', 'upload.xlsx')[:200],
                                       rows=data, created_by=actor.user)
    validate(actor, batch)
    return batch


def _decimal(value, errors, row, column, minimum=Decimal('0')):
    if value in ('', None):
        return None
    try:
        number = Decimal(str(value))
    except InvalidOperation:
        errors.append({'row': row, 'column': column, 'message': f'"{value}" is not a number.'})
        return None
    if number < minimum:
        errors.append({'row': row, 'column': column, 'message': f'Must be at least {minimum}.'})
    return number


def validate(actor, batch):
    errors = VALIDATORS[batch.import_type](actor, batch.rows)
    batch.errors = errors
    batch.status = 'FAILED' if errors else 'VALIDATED'
    batch.save(update_fields=['errors', 'status', 'updated_at'])
    return batch


def _validate_opening_inventory(actor, rows):
    errors, seen_barcodes, seen_huids, seen_serials = [], set(), set(), set()
    permitted = set(allowed_locations(actor.tenant, actor.user, 'adjust').values_list('code', flat=True))
    for number, row in enumerate(rows, 2):
        sku = SKU.objects.filter(tenant=actor.tenant, code=row.get('sku_code')).select_related('item', 'location').first()
        if sku is None:
            errors.append({'row': number, 'column': 'sku_code', 'message': f'SKU {row.get("sku_code")} not found.'})
            continue
        if row.get('location_code') != sku.location.code:
            errors.append({'row': number, 'column': 'location_code', 'message': f'SKU {sku.code} belongs to {sku.location.code}.'})
        if sku.location.code not in permitted:
            errors.append({'row': number, 'column': 'location_code', 'message': 'You cannot post opening stock at this location.'})
        if row.get('bin_code') and not Bin.objects.filter(location=sku.location, code=row['bin_code']).exists():
            errors.append({'row': number, 'column': 'bin_code', 'message': f'Bin {row["bin_code"]} not found in {sku.location.code}.'})
        quantity = _decimal(row.get('quantity') or '1', errors, number, 'quantity', Decimal('0.001'))
        _decimal(row.get('unit_cost'), errors, number, 'unit_cost')
        weights = [_decimal(row.get(c), errors, number, c) for c in ('gross_weight', 'stone_weight', 'other_weight')]
        if sku.item.serial_tracking:
            barcode, serial, huid = row.get('barcode'), row.get('serial_no') or row.get('barcode'), (row.get('huid') or '').upper()
            if not barcode:
                errors.append({'row': number, 'column': 'barcode', 'message': 'Serialized jewellery needs a barcode per piece.'})
            if quantity not in (None, Decimal('1')):
                errors.append({'row': number, 'column': 'quantity', 'message': 'Each serialized piece is one row with quantity 1.'})
            for value, seen, column, lookup in ((barcode, seen_barcodes, 'barcode', 'barcode'), (serial, seen_serials, 'serial_no', 'serial_no'),
                                                (huid, seen_huids, 'huid', 'huid')):
                if not value:
                    continue
                if value in seen:
                    errors.append({'row': number, 'column': column, 'message': f'Duplicate {column} {value} in this file.'})
                elif JewelleryUnit.objects.filter(tenant=actor.tenant, **{lookup: value}).exists():
                    errors.append({'row': number, 'column': column, 'message': f'{column} {value} already exists.'})
                seen.add(value)
            gross, stone, other = (w or Decimal('0') for w in weights)
            if stone + other > gross and gross:
                errors.append({'row': number, 'column': 'gross_weight', 'message': 'Stone + other weight exceeds gross weight.'})
    return errors


def _validate_locations(actor, rows):
    errors, seen = [], set()
    types = dict(Location.TYPES)
    for number, row in enumerate(rows, 2):
        code = row.get('code')
        if code in seen or Location.objects.filter(tenant=actor.tenant, code=code).exists():
            errors.append({'row': number, 'column': 'code', 'message': f'Location {code} already exists.'})
        seen.add(code)
        if (row.get('location_type') or '').upper() not in types:
            errors.append({'row': number, 'column': 'location_type', 'message': f'Use one of {", ".join(types)}.'})
        if row.get('gstin') and len(row['gstin']) != 15:
            errors.append({'row': number, 'column': 'gstin', 'message': 'GSTIN must be 15 characters.'})
    return errors


def _validate_transfer_orders(actor, rows):
    errors = []
    permitted = set(allowed_locations(actor.tenant, actor.user, 'create_transfer').values_list('code', flat=True))
    for number, row in enumerate(rows, 2):
        for column in ('from_location', 'to_location'):
            if not Location.objects.filter(tenant=actor.tenant, code=row.get(column)).exists():
                errors.append({'row': number, 'column': column, 'message': f'Location {row.get(column)} not found.'})
        if row.get('from_location') not in permitted:
            errors.append({'row': number, 'column': 'from_location', 'message': 'You cannot create transfers from this location.'})
        if row.get('barcode'):
            unit = find_unit(actor.tenant, row['barcode'])
            if unit is None or unit.current_location is None or unit.current_location.code != row.get('from_location'):
                errors.append({'row': number, 'column': 'barcode', 'message': f'{row["barcode"]} is not in stock at {row.get("from_location")}.'})
        elif not SKU.objects.filter(tenant=actor.tenant, code=row.get('sku_code'), location__code=row.get('from_location')).exists():
            errors.append({'row': number, 'column': 'sku_code', 'message': f'SKU {row.get("sku_code")} not found at {row.get("from_location")}.'})
        else:
            _decimal(row.get('quantity'), errors, number, 'quantity', Decimal('0.001'))
    return errors


@transaction.atomic
def run_import(actor, batch):
    batch = ImportBatch.objects.select_for_update().get(pk=batch.pk, tenant=actor.tenant)
    if batch.status != 'VALIDATED':
        raise InventoryError('Only a validated batch without errors can be imported.')
    validate(actor, batch)
    if batch.errors:
        raise InventoryError('The data changed since validation - review the error report.')
    result = IMPORTERS[batch.import_type](actor, batch.rows)
    batch.status, batch.result = 'IMPORTED', result
    batch.save(update_fields=['status', 'result', 'updated_at'])
    return batch


def _import_opening_inventory(actor, rows):
    units = pieces = 0
    for row in rows:
        sku = SKU.objects.select_related('item', 'location').get(tenant=actor.tenant, code=row['sku_code'])
        bin_ = Bin.objects.filter(location=sku.location, code=row['bin_code']).first() if row.get('bin_code') else None
        cost = Decimal(row['unit_cost']) if row.get('unit_cost') else None
        if sku.item.serial_tracking:
            weights = {k: Decimal(row[k]) for k in ('gross_weight', 'stone_weight', 'other_weight') if row.get(k)}
            unit = register_unit(actor, sku=sku, barcode=row['barcode'], serial_no=row.get('serial_no') or row['barcode'],
                                 huid=row.get('huid'), purity=row.get('purity') or None, certificate_no=row.get('certificate_no') or None,
                                 metal_cost=cost, **weights)
            post_receipt(actor, location=sku.location, sku=sku, unit=unit, bin=bin_, unit_cost=cost)
            units += 1
        else:
            quantity = Decimal(row.get('quantity') or '1')
            gross = Decimal(row['gross_weight']) if row.get('gross_weight') else None
            post_receipt(actor, location=sku.location, sku=sku, quantity=quantity, unit_cost=cost, bin=bin_, gross_weight=gross,
                         net_weight=gross - Decimal(row.get('stone_weight') or 0) - Decimal(row.get('other_weight') or 0) if gross is not None else None)
            pieces += quantity
    return {'jewellery_units': units, 'quantity': str(pieces), 'rows': len(rows)}


def _import_locations(actor, rows):
    for row in rows:
        Location.objects.create(tenant=actor.tenant, code=row['code'], name=row['name'], location_type=row['location_type'].upper(),
                                city=row.get('city', ''), state=row.get('state', ''), pin=row.get('pin', ''), gstin=row.get('gstin', ''),
                                is_warehouse=row.get('is_warehouse', '').lower() in ('1', 'yes', 'true', 'y'),
                                allow_pos=row.get('allow_pos', '').lower() in ('1', 'yes', 'true', 'y'), created_by=actor.user)
    return {'locations': len(rows)}


def _import_transfer_orders(actor, rows):
    groups = {}
    for row in rows:
        groups.setdefault((row['transfer_ref'], row['from_location'], row['to_location']), []).append(row)
    created = []
    for (ref, source, destination), group in groups.items():
        from_location = Location.objects.get(tenant=actor.tenant, code=source)
        to_location = Location.objects.get(tenant=actor.tenant, code=destination)
        lines = []
        for row in group:
            if row.get('barcode'):
                lines.append({'unit': find_unit(actor.tenant, row['barcode'])})
            else:
                lines.append({'sku': SKU.objects.get(tenant=actor.tenant, code=row['sku_code'], location=from_location),
                              'quantity': Decimal(row['quantity']),
                              'from_bin': Bin.objects.filter(location=from_location, code=row.get('from_bin')).first() if row.get('from_bin') else None,
                              'to_bin': Bin.objects.filter(location=to_location, code=row.get('to_bin')).first() if row.get('to_bin') else None})
        order = create_transfer_order(actor, from_location=from_location, to_location=to_location, lines=lines, source_type='IMPORT',
                                      remarks=f'Excel import {ref}')
        created.append(order.transfer_no)
    return {'transfer_orders': created}


VALIDATORS = {'opening_inventory': _validate_opening_inventory, 'locations': _validate_locations,
              'transfer_orders': _validate_transfer_orders}
IMPORTERS = {'opening_inventory': _import_opening_inventory, 'locations': _import_locations,
             'transfer_orders': _import_transfer_orders}
