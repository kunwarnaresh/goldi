from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile
from xml.sax.saxutils import escape

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
except ModuleNotFoundError:
    Workbook = None


TEMPLATE_FIELDS = {
    'customers': [('name', 'Customer name', True), ('phone', 'Phone', False), ('email', 'Email', False), ('gstin', 'GSTIN', False), ('customer_type', 'Customer type', True), ('is_active', 'Active', False)],
    'suppliers': [('name', 'Vendor name', True), ('phone', 'Phone', False), ('email', 'Email', False), ('gstin', 'GSTIN', False), ('pan', 'PAN', False), ('is_active', 'Active', False)],
    'products': [('sku', 'SKU', True), ('name', 'Item name', True), ('metal_type', 'Metal type', True), ('purity', 'Purity', False), ('purchase_price', 'Purchase price', False), ('sale_price', 'Sale price', False), ('barcode', 'Barcode', False), ('hsn_code', 'HSN code', False), ('is_active', 'Active', False)],
    'retail-items': [('item_number', 'Item number', True), ('name', 'Item name', True), ('company', 'Company ID', True), ('division', 'Division ID', True), ('special_group', 'Special group ID', True), ('category', 'Category ID', True), ('subcategory', 'SubCategory ID', True), ('base_uom', 'Base UOM ID', True), ('status', 'Status', True)],
    'divisions': [('code', 'Division code', True), ('name', 'Division name', True), ('company', 'Company ID', True), ('active', 'Active', False), ('blocked', 'Blocked', False)],
    'special-groups': [('code', 'Special group code', True), ('name', 'Special group name', True), ('division', 'Division ID', True), ('active', 'Active', False), ('blocked', 'Blocked', False)],
    'categories': [('code', 'Category code', True), ('name', 'Category name', True), ('special_group', 'Special group ID', True), ('parent', 'Parent category ID', False), ('active', 'Active', False), ('blocked', 'Blocked', False)],
    'subcategories': [('code', 'SubCategory code', True), ('name', 'SubCategory name', True), ('category', 'Category ID', True), ('active', 'Active', False), ('blocked', 'Blocked', False)],
    'units-of-measure': [('code', 'UOM code', True), ('name', 'UOM name', True), ('company', 'Company ID', False), ('is_active', 'Active', False)],
    'item-variants': [('code', 'Variant code', True), ('name', 'Variant name', True), ('item', 'Item ID', True), ('color', 'Color', False), ('size', 'Size', False), ('material', 'Material', False), ('active', 'Active', False), ('blocked', 'Blocked', False)],
    'skus': [('code', 'SKU code', True), ('item', 'Item ID', True), ('variant', 'Variant ID', False), ('location', 'Location ID', False), ('cost', 'Cost', False), ('price', 'Price', False), ('active', 'Active', False), ('blocked', 'Blocked', False)],
    'retail-barcodes': [('value', 'Barcode', True), ('barcode_type', 'Barcode type', True), ('sku', 'SKU ID', True), ('uom', 'UOM ID', True), ('quantity', 'Quantity', False), ('is_primary', 'Primary barcode', False), ('active', 'Active', False)],
    'item-uoms': [('item', 'Item ID', True), ('uom', 'UOM ID', True), ('quantity_per_uom', 'Quantity per UOM', True), ('is_inventory_uom', 'Inventory UOM', False), ('is_purchase_uom', 'Purchase UOM', False), ('is_sales_uom', 'Sales UOM', False)],
    'item-locations': [('item', 'Item ID', True), ('variant', 'Variant ID', False), ('location', 'Location ID', True), ('default_bin', 'Default bin ID', False), ('reorder_point', 'Reorder point', False), ('reorder_quantity', 'Reorder quantity', False), ('active', 'Active', False)],
    'warehouses': [('code', 'Warehouse code', True), ('name', 'Warehouse name', True), ('company', 'Company ID', False), ('location_type', 'Location type', True), ('status', 'Status', True)],
    'locations': [('location_code', 'Location code', True), ('location_name', 'Location name', True), ('company', 'Company ID', True), ('warehouse', 'Warehouse ID', False), ('status', 'Status', True)],
    'zones': [('code', 'Zone code', True), ('name', 'Zone name', True), ('warehouse', 'Warehouse ID', True), ('zone_type', 'Zone type', True), ('priority', 'Priority', False), ('status', 'Status', True)],
    'bins': [('code', 'Bin code', True), ('warehouse', 'Warehouse ID', True), ('zone', 'Zone ID', False), ('bin_type', 'Bin type', True), ('capacity', 'Capacity', False), ('is_active', 'Active', False)],
    'gst-rates': [('code', 'Rate code', True), ('description', 'Description', True), ('effective_from', 'Effective from YYYY-MM-DD', True), ('cgst_rate', 'CGST %', False), ('sgst_rate', 'SGST %', False), ('igst_rate', 'IGST %', False), ('cess_rate', 'Cess %', False), ('status', 'Status', True)],
    'finance-vouchers': [('voucher_no', 'Voucher number', False), ('company', 'Company ID', True), ('batch', 'Batch ID', True), ('voucher_type', 'Voucher type ID', True), ('document_no', 'Document number', False), ('voucher_date', 'Voucher date YYYY-MM-DD', True), ('narration', 'Narration', False)],
}


def template_fields(slug, model_fields):
    return TEMPLATE_FIELDS.get(slug, [(field, field.replace('_', ' ').title(), field not in {'id', 'created_at', 'updated_at'}) for field in model_fields if field not in {'id', 'created_at', 'updated_at'}])


def build_template(slug, label, model_fields):
    fields = template_fields(slug, model_fields)
    if Workbook is None:
        return build_minimal_xlsx(label, fields)
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = '01_Instructions'
    instructions.append([f'{label} import template'])
    instructions.append(['Purpose', f'Create or update {label} records through the ERP import staging process.'])
    instructions.append(['Process', 'Fill 02_Data, save as XLSX, upload, validate, correct errors, then approve/import.'])
    instructions.append(['Rules', 'Do not rename columns, add formulas, or delete required fields. IDs must reference existing ERP records.'])
    instructions.append(['Import mode', 'Choose Create, Update, or Create + Update in the import workflow.'])
    data = workbook.create_sheet('02_Data')
    headers = [field_name + (' *' if required else '') for field_name, _, required in fields]
    data.append(headers)
    examples = workbook.create_sheet('03_Examples')
    examples.append(headers)
    examples.append([f'EXAMPLE-{index + 1}' if index == 0 else ('Example value' if required else '') for index, (_, _, required) in enumerate(fields)])
    definitions = workbook.create_sheet('04_Field_Definitions')
    definitions.append(['Field name', 'Description', 'Required', 'Data type', 'Example'])
    for field_name, description, required in fields:
        definitions.append([field_name, description, 'Yes' if required else 'No', 'Text', f'Enter {description.lower()}'])
    lookups = workbook.create_sheet('05_Lookup_Values')
    lookups.append(['Lookup', 'Allowed value'])
    lookups.append(['Active', 'TRUE'])
    lookups.append(['Active', 'FALSE'])
    errors = workbook.create_sheet('06_Errors')
    errors.append(['Row number', 'Column', 'Entered value', 'Error type', 'Error message', 'Suggested correction'])
    header_fill = PatternFill('solid', fgColor='1F4E78')
    for sheet in workbook.worksheets:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        sheet.column_dimensions['A'].width = 24
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = header_fill
    data.column_dimensions['A'].width = 24
    active_validation = DataValidation(type='list', formula1='"TRUE,FALSE"', allow_blank=True)
    data.add_data_validation(active_validation)
    for index, (field_name, _, _) in enumerate(fields, start=1):
        if field_name in {'active', 'is_active', 'blocked', 'is_primary', 'is_inventory_uom', 'is_purchase_uom', 'is_sales_uom'}:
            active_validation.add(data.cell(row=2, column=index))
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output.getvalue()


def build_minimal_xlsx(label, fields):
    sheets = {
        '01_Instructions': [[f'{label} import template'], ['Purpose', f'Create or update {label} records through ERP staging.'], ['Process', 'Fill 02_Data, upload, validate, correct errors, then approve/import.']],
        '02_Data': [[field + (' *' if required else '') for field, _, required in fields]],
        '03_Examples': [[field + (' *' if required else '') for field, _, required in fields], ['Example' for _ in fields]],
        '04_Field_Definitions': [['Field name', 'Description', 'Required', 'Data type', 'Example']] + [[field, description, 'Yes' if required else 'No', 'Text', f'Enter {description.lower()}'] for field, description, required in fields],
        '05_Lookup_Values': [['Lookup', 'Allowed value'], ['Active', 'TRUE'], ['Active', 'FALSE']],
        '06_Errors': [['Row number', 'Column', 'Entered value', 'Error type', 'Error message', 'Suggested correction']],
    }
    content_types = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>', '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">', '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>', '<Default Extension="xml" ContentType="application/xml"/>', '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>']
    relationships = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>', '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">', '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>', '</Relationships>']
    workbook = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>', '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets>']
    workbook_rels = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>', '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">']
    for index, sheet_name in enumerate(sheets, 1):
        workbook.append(f'<sheet name="{escape(sheet_name)}" sheetId="{index}" r:id="rId{index}"/>')
        workbook_rels.append(f'<Relationship Id="rId{index}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet{index}.xml"/>')
        content_types.append(f'<Override PartName="/xl/worksheets/sheet{index}.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>')
    workbook.extend(['</sheets></workbook>'])
    workbook_rels.append('</Relationships>')
    content_types.append('</Types>')
    output = BytesIO()
    with ZipFile(output, 'w', ZIP_DEFLATED) as archive:
        archive.writestr('[Content_Types].xml', ''.join(content_types))
        archive.writestr('_rels/.rels', ''.join(relationships))
        archive.writestr('xl/workbook.xml', ''.join(workbook))
        archive.writestr('xl/_rels/workbook.xml.rels', ''.join(workbook_rels))
        for index, rows in enumerate(sheets.values(), 1):
            xml = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>', '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>']
            for row_number, row in enumerate(rows, 1):
                xml.append(f'<row r="{row_number}">')
                for column_number, value in enumerate(row, 1):
                    column = chr(64 + min(column_number, 26))
                    xml.append(f'<c r="{column}{row_number}" t="inlineStr"><is><t>{escape(str(value))}</t></is></c>')
                xml.append('</row>')
            xml.append('</sheetData></worksheet>')
            archive.writestr(f'xl/worksheets/sheet{index}.xml', ''.join(xml))
    return output.getvalue()
