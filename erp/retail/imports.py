"""Excel template / export / staged import for the store & POS masters.

Upload -> validate (a full dry run through the same forms the UI uses, rolled back) -> preview with per-row errors ->
download error report / re-upload corrected file -> import (all rows or none)."""
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.contrib.auth.hashers import make_password
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from erp.models import (
    Company, Location, POSRole, POSStaff, POSStaffAssignment, POSTerminal, RetailImportBatch, Store, StoreTender, Tender,
)

from . import forms as f
from . import services

HASH_PREFIX = 'hashed$'

# column -> getter used for export. Column order is the template order.
COLUMNS = {
    'locations': [
        ('company_code', lambda o: o.company.company_code), ('location_code', lambda o: o.location_code),
        ('location_name', lambda o: o.location_name), ('branch_code', lambda o: o.branch.branch_code if o.branch else ''),
        ('address', lambda o: o.address), ('city', lambda o: o.city), ('state_code', lambda o: o.state_code),
        ('country', lambda o: o.country), ('status', lambda o: o.status),
    ],
    'stores': [
        ('company_code', lambda o: o.company.company_code), ('store_code', lambda o: o.code), ('store_name', lambda o: o.name),
        ('store_type', lambda o: o.store_type), ('location_code', lambda o: o.location.location_code if o.location else ''),
        ('address', lambda o: o.address), ('city', lambda o: o.city), ('state', lambda o: o.state),
        ('country', lambda o: o.country), ('pin_code', lambda o: o.pin_code), ('gstin', lambda o: o.gstin),
        ('phone', lambda o: o.phone), ('email', lambda o: o.email), ('status', lambda o: o.status),
        ('opening_date', lambda o: o.opening_date),
    ],
    'staff': [
        ('employee_code', lambda o: o.employee_code), ('staff_name', lambda o: o.name), ('first_name', lambda o: o.first_name),
        ('last_name', lambda o: o.last_name), ('mobile', lambda o: o.mobile), ('email', lambda o: o.email),
        ('designation', lambda o: o.designation), ('role_code', lambda o: o.staff_role.code if o.staff_role else ''),
        ('company_code', lambda o: o.store.company.company_code), ('store_code', lambda o: o.store.code),
        ('location_code', lambda o: o.store.location.location_code if o.store.location else ''),
        ('login_id', lambda o: o.login_id or ''), ('password', lambda o: ''),  # never exported
        ('pos_login_enabled', lambda o: o.pos_access), ('active', lambda o: o.is_active),
        ('joining_date', lambda o: o.joining_date),
    ],
    'terminals': [
        ('company_code', lambda o: o.store.company.company_code), ('store_code', lambda o: o.store.code),
        ('terminal_code', lambda o: o.code), ('terminal_name', lambda o: o.name), ('terminal_type', lambda o: o.terminal_type),
        ('device_id', lambda o: o.device_id), ('serial_number', lambda o: o.serial_number),
        ('ip_address', lambda o: o.ip_address or ''), ('mac_address', lambda o: o.mac_address),
        ('receipt_printer', lambda o: o.receipt_printer), ('status', lambda o: o.status),
    ],
    'tenders': [
        ('tender_code', lambda o: o.code), ('tender_name', lambda o: o.name), ('tender_type', lambda o: o.tender_type),
        ('requires_reference', lambda o: o.requires_reference), ('requires_approval', lambda o: o.requires_approval),
        ('allow_refund', lambda o: o.allow_refund), ('allow_change', lambda o: o.allow_change),
        ('allow_split_payment', lambda o: o.allow_split_payment), ('allow_partial_payment', lambda o: o.allow_partial_payment),
        ('minimum_amount', lambda o: o.minimum_amount), ('maximum_amount', lambda o: o.maximum_amount),
        ('status', lambda o: o.status),
    ],
    'store-tenders': [
        ('company_code', lambda o: o.store.company.company_code), ('store_code', lambda o: o.store.code),
        ('tender_code', lambda o: o.tender.code), ('active', lambda o: o.active), ('default', lambda o: o.is_default),
        ('sequence', lambda o: o.sequence), ('requires_reference', lambda o: o.requires_reference),
        ('allow_refund', lambda o: o.allow_refund), ('allow_change', lambda o: o.allow_change),
        ('allow_split_payment', lambda o: o.allow_split_payment),
    ],
    'staff-terminals': [
        ('employee_code', lambda o: o.staff.employee_code), ('company_code', lambda o: o.terminal.store.company.company_code),
        ('store_code', lambda o: o.terminal.store.code), ('terminal_code', lambda o: o.terminal.code),
        ('role', lambda o: o.role), ('active', lambda o: o.active), ('effective_from', lambda o: o.effective_from),
    ],
}
REQUIRED = {
    'locations': ['company_code', 'location_code', 'location_name'],
    'stores': ['company_code', 'store_code', 'store_name', 'location_code'],
    'staff': ['employee_code', 'staff_name', 'role_code', 'store_code'],
    'terminals': ['store_code', 'terminal_code', 'terminal_name'],
    'tenders': ['tender_code', 'tender_name', 'tender_type'],
    'store-tenders': ['store_code', 'tender_code'],
    'staff-terminals': ['employee_code', 'store_code', 'terminal_code'],
}
QUERYSETS = {
    'locations': lambda: Location.objects.select_related('company', 'branch').order_by('location_code'),
    'stores': lambda: Store.objects.select_related('company', 'location').order_by('code'),
    'staff': lambda: POSStaff.objects.select_related('store__company', 'store__location', 'staff_role').order_by('employee_code'),
    'terminals': lambda: POSTerminal.objects.select_related('store__company').order_by('store__code', 'code'),
    'tenders': lambda: Tender.objects.order_by('code'),
    'store-tenders': lambda: StoreTender.objects.select_related('store__company', 'tender').order_by('store__code', 'sequence'),
    'staff-terminals': lambda: POSStaffAssignment.objects.select_related('staff', 'terminal__store__company').order_by('staff__employee_code'),
}


class RowError(Exception):
    pass


# --- workbook helpers ------------------------------------------------------------------------------------------------

def _cell(value):
    if isinstance(value, bool):
        return 'Yes' if value else 'No'
    if isinstance(value, Decimal):
        return float(value)
    return value if value is not None else ''


def _workbook(headers, rows, title):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='DCE8F5')
    for row in rows:
        ws.append([_cell(v) for v in row])
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width = max(12, min(40, max(len(str(c.value or '')) for c in column) + 2))
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def template_workbook(kind):
    return _workbook([c for c, _ in COLUMNS[kind]], [], kind)


def export_workbook(kind, queryset=None):
    qs = queryset if queryset is not None else QUERYSETS[kind]()
    return _workbook([c for c, _ in COLUMNS[kind]], [[get(o) for _, get in COLUMNS[kind]] for o in qs], kind)


def error_workbook(batch):
    columns = [c for c, _ in COLUMNS[batch.import_type]]
    by_row = {}
    for error in batch.errors:
        by_row.setdefault(error['row'], []).append(error['message'])
    rows = [[row.get(c, '') if c != 'password' else '' for c in columns] + ['; '.join(by_row.get(n, []))]
            for n, row in enumerate(batch.rows, 2)]
    return _workbook(columns + ['errors'], rows, 'errors')


# --- parsing ---------------------------------------------------------------------------------------------------------

def _text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, (datetime, date)):
        return value.strftime('%Y-%m-%d')
    return str(value).strip()


def _bool(value, default=None):
    text = _text(value).lower()
    if text == '':
        return default
    if text in ('yes', 'y', 'true', '1', 'active', 'enabled'):
        return True
    if text in ('no', 'n', 'false', '0', 'inactive', 'disabled'):
        return False
    raise RowError(f'"{value}" is not a yes/no value.')


def _decimal(value):
    text = _text(value)
    if not text:
        return ''
    try:
        return str(Decimal(text))
    except InvalidOperation as exc:
        raise RowError(f'"{value}" is not a number.') from exc


def read_rows(uploaded, kind):
    """Read an uploaded workbook into plain dict rows. Plain-text passwords are hashed immediately."""
    wb = load_workbook(uploaded, read_only=True, data_only=True)
    ws = wb.active
    raw = list(ws.iter_rows(values_only=True))
    if not raw:
        return [], [{'row': 1, 'column': '', 'message': 'The workbook is empty.'}]
    headers = [_text(h).lower() for h in raw[0]]
    expected = [c for c, _ in COLUMNS[kind]]
    errors = []
    missing = [c for c in REQUIRED[kind] if c not in headers]
    if missing:
        errors.append({'row': 1, 'column': '', 'message': f'Missing column(s): {", ".join(missing)}.'})
    rows = []
    for values in raw[1:]:
        if values is None or all(v in (None, '') for v in values):
            continue
        row = {h: _text(v) for h, v in zip(headers, values) if h in expected}
        if kind == 'staff' and row.get('password'):
            password = row['password']
            row['password'] = HASH_PREFIX + make_password(password) if len(password) >= services.PASSWORD_MIN_LENGTH else '!too-short'
        rows.append(row)
    return rows, errors


# --- row handlers: each maps a row onto the same ModelForm the UI uses ------------------------------------------------

def _company(row):
    code = row.get('company_code')
    if not code:
        return None
    company = Company.objects.filter(company_code__iexact=code).first()
    if company is None:
        raise RowError(f'Company {code} not found.')
    return company


def _store(row):
    qs = Store.objects.filter(code__iexact=row.get('store_code', ''))
    company = _company(row)
    if company:
        qs = qs.filter(company=company)
    stores = list(qs[:2])
    if not stores:
        raise RowError(f'Store {row.get("store_code")} not found.')
    if len(stores) > 1:
        raise RowError(f'Store code {row.get("store_code")} exists in several companies; fill company_code.')
    return stores[0]


def _form_errors(form):
    return [f'{field if field != "__all__" else "row"}: {" ".join(msgs)}' for field, msgs in form.errors.items()]


def _save(form, instance_exists):
    if not form.is_valid():
        raise RowError(' | '.join(_form_errors(form)))
    return form, instance_exists


def handle_locations(row, request):
    company = _company(row)
    if company is None:
        raise RowError('company_code is required.')
    instance = Location.objects.filter(company=company, location_code__iexact=row['location_code']).first()
    branch = None
    if row.get('branch_code'):
        branch = company.branches.filter(branch_code__iexact=row['branch_code']).first()
        if branch is None:
            raise RowError(f'Branch {row["branch_code"]} not found.')
    data = {'company': company.pk, 'branch': branch.pk if branch else '', 'location_code': row['location_code'],
            'location_name': row['location_name'], 'address': row.get('address', ''), 'city': row.get('city', ''),
            'state_code': row.get('state_code', ''), 'country': row.get('country') or 'IN',
            'status': (row.get('status') or 'active').lower()}
    form = f.LocationForm(data, instance=instance)
    _save(form, instance)
    services.save_with_audit(form.save(commit=False), ['location_code', 'location_name', 'status'], request, 'Location')
    return instance is not None


def handle_stores(row, request):
    company = _company(row)
    if company is None:
        raise RowError('company_code is required.')
    instance = Store.objects.filter(company=company, code__iexact=row['store_code']).first()
    location = Location.objects.filter(company=company, location_code__iexact=row['location_code']).first()
    if location is None:
        raise RowError(f'Location {row["location_code"]} not found in company {company.company_code}.')
    data = {'company': company.pk, 'code': row['store_code'], 'name': row['store_name'],
            'store_type': (row.get('store_type') or 'showroom').lower(), 'location': location.pk,
            'address': row.get('address', ''), 'city': row.get('city', ''), 'state': row.get('state', ''),
            'country': row.get('country') or 'India', 'pin_code': row.get('pin_code', ''), 'gstin': row.get('gstin', ''),
            'phone': row.get('phone', ''), 'email': row.get('email', ''), 'status': (row.get('status') or 'active').lower(),
            'opening_date': row.get('opening_date', ''), 'manager': instance.manager_id if instance and instance.manager_id else ''}
    form = f.StoreForm(data, instance=instance)
    _save(form, instance)
    services.save_with_audit(form.save(commit=False), services.STORE_AUDIT_FIELDS, request, 'Store')
    return instance is not None


class _ImportStaffForm(f.StaffForm):
    def __init__(self, *args, password_hash=None, **kwargs):
        self.password_hash = password_hash
        super().__init__(*args, **kwargs)

    def clean(self):
        data = super().clean()
        if self.password_hash and 'password' in self.errors:
            del self.errors['password']
        return data


def handle_staff(row, request):
    instance = POSStaff.objects.filter(employee_code__iexact=row['employee_code']).first()
    role = POSRole.objects.filter(code__iexact=row['role_code']).first()
    if role is None:
        raise RowError(f'Role {row["role_code"]} not found.')
    store = _store(row)
    location = None
    if row.get('location_code'):
        location = Location.objects.filter(company=store.company, location_code__iexact=row['location_code']).first()
        if location is None:
            raise RowError(f'Location {row["location_code"]} not found.')
    password = row.get('password', '')
    if password == '!too-short':
        raise RowError(f'password: must be at least {services.PASSWORD_MIN_LENGTH} characters.')
    password_hash = password[len(HASH_PREFIX):] if password.startswith(HASH_PREFIX) else None
    pos_access = _bool(row.get('pos_login_enabled'), True)
    data = {'employee_code': row['employee_code'], 'name': row['staff_name'], 'first_name': row.get('first_name', ''),
            'last_name': row.get('last_name', ''), 'mobile': row.get('mobile', ''), 'email': row.get('email', ''),
            'designation': row.get('designation', ''), 'staff_role': role.pk, 'store': store.pk,
            'location': location.pk if location else '', 'login_id': row.get('login_id', ''),
            'joining_date': row.get('joining_date', ''),
            'default_terminal': instance.default_terminal_id if instance and instance.default_terminal_id and instance.store_id == store.pk else ''}
    if pos_access:
        data['pos_access'] = 'on'
    if _bool(row.get('active'), True):
        data['is_active'] = 'on'
    form = _ImportStaffForm(data, instance=instance, password_hash=password_hash)
    _save(form, instance)
    staff = form.save(commit=False)
    if password_hash:
        staff.pin_hash = password_hash
        staff.password_changed_at = timezone.now()
    services.save_staff(staff, request, location=form.cleaned_data.get('location'))
    return instance is not None


def handle_terminals(row, request):
    store = _store(row)
    instance = POSTerminal.objects.filter(store=store, code__iexact=row['terminal_code']).first()
    data = {'store': store.pk, 'code': row['terminal_code'], 'name': row['terminal_name'],
            'terminal_type': (row.get('terminal_type') or 'counter').lower(), 'device_id': row.get('device_id', ''),
            'serial_number': row.get('serial_number', ''), 'ip_address': row.get('ip_address', ''),
            'mac_address': row.get('mac_address', ''), 'receipt_printer': row.get('receipt_printer', ''),
            'status': (row.get('status') or 'active').lower(),
            'default_tender': instance.default_tender_id if instance and instance.default_tender_id else ''}
    form = f.TerminalForm(data, instance=instance)
    _save(form, instance)
    services.save_with_audit(form.save(commit=False), services.TERMINAL_AUDIT_FIELDS, request, 'POS terminal')
    return instance is not None


def handle_tenders(row, request):
    instance = Tender.objects.filter(code__iexact=row['tender_code']).first()
    data = {'code': row['tender_code'], 'name': row['tender_name'], 'tender_type': row['tender_type'].lower(),
            'status': (row.get('status') or 'active').lower(), 'currency': 'INR',
            'minimum_amount': _decimal(row.get('minimum_amount')), 'maximum_amount': _decimal(row.get('maximum_amount'))}
    defaults = {'requires_reference': False, 'requires_approval': False, 'allow_refund': True, 'allow_change': False,
                'allow_split_payment': True, 'allow_partial_payment': True}
    for flag, default in defaults.items():
        if _bool(row.get(flag), default):
            data[flag] = 'on'
    form = f.TenderForm(data, instance=instance)
    _save(form, instance)
    services.save_with_audit(form.save(commit=False), services.TENDER_AUDIT_FIELDS, request, 'Tender')
    return instance is not None


def handle_store_tenders(row, request):
    store = _store(row)
    tender = Tender.objects.filter(code__iexact=row['tender_code']).first()
    if tender is None:
        raise RowError(f'Tender {row["tender_code"]} not found.')
    instance = StoreTender.objects.filter(store=store, tender=tender).first()
    data = {'tender': tender.pk, 'sequence': row.get('sequence') or 10}
    defaults = {'active': True, 'default': False, 'requires_reference': tender.requires_reference,
                'allow_refund': tender.allow_refund, 'allow_change': tender.allow_change,
                'allow_split_payment': tender.allow_split_payment}
    for flag, default in defaults.items():
        if _bool(row.get(flag), default):
            data['is_default' if flag == 'default' else flag] = 'on'
    form = f.StoreTenderForm(data, instance=instance, store=store)
    _save(form, instance)
    services.assign_store_tender(form.save(commit=False), request)
    return instance is not None


def handle_staff_terminals(row, request):
    staff = POSStaff.objects.filter(employee_code__iexact=row['employee_code']).first()
    if staff is None:
        raise RowError(f'Staff {row["employee_code"]} not found.')
    store = _store(row)
    terminal = POSTerminal.objects.filter(store=store, code__iexact=row['terminal_code']).first()
    if terminal is None:
        raise RowError(f'Terminal {row["terminal_code"]} not found in store {store.code}.')
    role = (row.get('role') or staff.role).lower()
    instance = POSStaffAssignment.objects.filter(staff=staff, terminal=terminal, role=role).first()
    data = {'terminal': terminal.pk, 'role': role, 'effective_from': row.get('effective_from') or timezone.localdate().isoformat()}
    if _bool(row.get('active'), True):
        data['active'] = 'on'
    form = f.AssignmentForm(data, instance=instance, staff=staff)
    _save(form, instance)
    assignment = form.save()
    services.audit(request, staff, 'assign', f'{staff.employee_code} assigned to terminal {terminal.code}.',
                   new={'terminal': terminal.code, 'role': role, 'active': assignment.active})
    return instance is not None


HANDLERS = {
    'locations': handle_locations, 'stores': handle_stores, 'staff': handle_staff, 'terminals': handle_terminals,
    'tenders': handle_tenders, 'store-tenders': handle_store_tenders, 'staff-terminals': handle_staff_terminals,
}


class _DryRun(Exception):
    pass


def _process(batch, request, commit):
    errors, created, updated = [], 0, 0
    handler = HANDLERS[batch.import_type]
    try:
        with transaction.atomic():
            for number, row in enumerate(batch.rows, 2):
                missing = [c for c in REQUIRED[batch.import_type] if not row.get(c)]
                if missing:
                    errors.append({'row': number, 'column': missing[0], 'message': f'Required: {", ".join(missing)}.'})
                    continue
                try:
                    with transaction.atomic():
                        existed = handler(row, request)
                    updated += existed
                    created += not existed
                except RowError as exc:
                    errors.append({'row': number, 'column': '', 'message': str(exc)})
                except Exception as exc:  # validation raised from model save, integrity errors, bad dates...
                    messages = getattr(exc, 'messages', None) or [str(exc)]
                    errors.append({'row': number, 'column': '', 'message': ' '.join(messages)})
            if errors or not commit:
                raise _DryRun
    except _DryRun:
        pass
    return errors, created, updated


def upload(uploaded, import_type, request):
    rows, errors = read_rows(uploaded, import_type)
    batch = RetailImportBatch.objects.create(import_type=import_type, file_name=getattr(uploaded, 'name', '')[:200],
                                             rows=rows, errors=errors, status='failed' if errors else 'validated',
                                             created_by=request.user if request.user.is_authenticated else None)
    if not errors:
        validate(batch, request)
    return batch


def validate(batch, request):
    errors, created, updated = _process(batch, request, commit=False)
    batch.errors = errors
    batch.status = 'failed' if errors else 'validated'
    batch.created_count, batch.updated_count = created, updated
    batch.save(update_fields=['errors', 'status', 'created_count', 'updated_count'])
    return batch


def run_import(batch, request):
    if batch.status != 'validated':
        raise ValueError('Only a batch that validated without errors can be imported.')
    errors, created, updated = _process(batch, request, commit=True)
    if errors:
        batch.errors, batch.status = errors, 'failed'
        batch.save(update_fields=['errors', 'status'])
        return batch
    # Hashes are no longer needed once imported.
    batch.rows = [{k: ('' if k == 'password' else v) for k, v in row.items()} for row in batch.rows]
    batch.status, batch.created_count, batch.updated_count, batch.imported_at = 'imported', created, updated, timezone.now()
    batch.save(update_fields=['rows', 'status', 'created_count', 'updated_count', 'imported_at'])
    services.audit(request, batch, 'import', f'Imported {batch.get_import_type_display()}: {created} created, {updated} updated.')
    return batch
